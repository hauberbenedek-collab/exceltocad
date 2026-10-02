"""Fogaskerék paraméterek beolvasása és ellenőrzése Excel fájlból."""

import json
import math
import sys
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook

KOTELEZO_OSZLOPOK = [
    "nev",
    "modul",
    "fogszam",
    "nyomaszog",
    "szelesseg",
    "furat_nevleges",
    "furat_felso",
    "furat_also",
    "turesmod",
]
TURESMODOK = ("nevleges", "kozep", "min", "max")


class EllenorzesiHiba(Exception):
    """Hibás bemeneti adat esetén váltódik ki."""


@dataclass
class FogaskerekAdat:
    nev: str
    modul: float
    fogszam: int
    nyomaszog: float
    szelesseg: float
    furat_nevleges: float
    furat_felso: float
    furat_also: float
    turesmod: str

    @property
    def osztokor_atmero(self):
        return self.modul * self.fogszam

    @property
    def fejkor_atmero(self):
        return self.osztokor_atmero + 2 * self.modul

    @property
    def labkor_atmero(self):
        return self.osztokor_atmero - 2.5 * self.modul

    @property
    def alapkor_atmero(self):
        return self.osztokor_atmero * math.cos(math.radians(self.nyomaszog))

    @property
    def furat_atmero(self):
        """A furat átmérője a kiválasztott tűrésmód szerint."""
        also = self.furat_nevleges + self.furat_also
        felso = self.furat_nevleges + self.furat_felso
        if self.turesmod == "min":
            return also
        if self.turesmod == "max":
            return felso
        if self.turesmod == "kozep":
            return (also + felso) / 2
        return self.furat_nevleges

    def szotarba(self):
        adat = asdict(self)
        adat.update(
            osztokor_atmero=round(self.osztokor_atmero, 4),
            fejkor_atmero=round(self.fejkor_atmero, 4),
            labkor_atmero=round(self.labkor_atmero, 4),
            alapkor_atmero=round(self.alapkor_atmero, 4),
            furat_atmero=round(self.furat_atmero, 4),
        )
        return adat


def _szamma(ertek, oszlop, sor):
    if isinstance(ertek, str):
        ertek = ertek.strip().replace(",", ".")
    try:
        return float(ertek)
    except (TypeError, ValueError):
        raise EllenorzesiHiba(
            f"{sor}. sor: a '{oszlop}' értéke szám kell legyen, de ez: {ertek!r}"
        )


def _egeszre(ertek, oszlop, sor):
    szam = _szamma(ertek, oszlop, sor)
    if not szam.is_integer():
        raise EllenorzesiHiba(
            f"{sor}. sor: a '{oszlop}' értéke egész szám kell legyen, de ez: {ertek!r}"
        )
    return int(szam)


def _ellenoriz(kerek, sor):
    if kerek.modul <= 0:
        raise EllenorzesiHiba(f"{sor}. sor: a 'modul' pozitív kell legyen")
    if kerek.fogszam < 8:
        raise EllenorzesiHiba(f"{sor}. sor: a 'fogszam' legalább 8 kell legyen")
    if not 10 <= kerek.nyomaszog <= 30:
        raise EllenorzesiHiba(
            f"{sor}. sor: a 'nyomaszog' 10 és 30 fok között kell legyen"
        )
    if kerek.szelesseg <= 0:
        raise EllenorzesiHiba(f"{sor}. sor: a 'szelesseg' pozitív kell legyen")
    if kerek.furat_nevleges <= 0:
        raise EllenorzesiHiba(f"{sor}. sor: a 'furat_nevleges' pozitív kell legyen")
    if kerek.furat_felso < kerek.furat_also:
        raise EllenorzesiHiba(
            f"{sor}. sor: a 'furat_felso' nem lehet kisebb, mint a 'furat_also'"
        )
    if kerek.turesmod not in TURESMODOK:
        raise EllenorzesiHiba(
            f"{sor}. sor: a 'turesmod' ezek egyike lehet: "
            f"{', '.join(TURESMODOK)}, de ez: {kerek.turesmod!r}"
        )
    if kerek.furat_nevleges + kerek.furat_felso >= kerek.labkor_atmero:
        raise EllenorzesiHiba(
            f"{sor}. sor: a furat túl nagy, kisebb kell legyen a lábkör "
            f"átmérőjénél ({kerek.labkor_atmero:.2f} mm)"
        )


def _sor_feldolgozasa(ertekek, sor):
    nev = ertekek["nev"]
    if nev is None or str(nev).strip() == "":
        raise EllenorzesiHiba(f"{sor}. sor: hiányzik a 'nev'")
    mod = ertekek["turesmod"]
    mod = "nevleges" if mod is None else str(mod).strip().lower()
    kerek = FogaskerekAdat(
        nev=str(nev).strip(),
        modul=_szamma(ertekek["modul"], "modul", sor),
        fogszam=_egeszre(ertekek["fogszam"], "fogszam", sor),
        nyomaszog=_szamma(ertekek["nyomaszog"], "nyomaszog", sor),
        szelesseg=_szamma(ertekek["szelesseg"], "szelesseg", sor),
        furat_nevleges=_szamma(ertekek["furat_nevleges"], "furat_nevleges", sor),
        furat_felso=_szamma(ertekek["furat_felso"], "furat_felso", sor),
        furat_also=_szamma(ertekek["furat_also"], "furat_also", sor),
        turesmod=mod,
    )
    _ellenoriz(kerek, sor)
    return kerek


def fogaskerekek_beolvasasa(utvonal, lap=None):
    """Beolvassa az összes fogaskereket az Excel fájlból (FogaskerekAdat lista)."""
    utvonal = Path(utvonal)
    if not utvonal.exists():
        raise FileNotFoundError(f"A bemeneti fájl nem található: {utvonal}")

    munkafuzet = load_workbook(utvonal, data_only=True, read_only=True)
    munkalap = munkafuzet[lap] if lap else munkafuzet.active
    sorok = munkalap.iter_rows(values_only=True)

    fejlec = next(sorok, None)
    if fejlec is None:
        raise EllenorzesiHiba("A munkalap üres")
    oszlopok = [str(c).strip().lower() if c is not None else "" for c in fejlec]
    hianyzo = [c for c in KOTELEZO_OSZLOPOK if c not in oszlopok]
    if hianyzo:
        raise EllenorzesiHiba("Hiányzó oszlop(ok): " + ", ".join(hianyzo))
    index = {nev: oszlopok.index(nev) for nev in KOTELEZO_OSZLOPOK}

    kerekek = []
    nevek = set()
    for sorszam, sor in enumerate(sorok, start=2):
        if sor is None or all(c is None or str(c).strip() == "" for c in sor):
            continue
        ertekek = {
            nev: (sor[i] if i < len(sor) else None) for nev, i in index.items()
        }
        kerek = _sor_feldolgozasa(ertekek, sorszam)
        if kerek.nev in nevek:
            raise EllenorzesiHiba(f"{sorszam}. sor: ismétlődő név: '{kerek.nev}'")
        nevek.add(kerek.nev)
        kerekek.append(kerek)
    munkafuzet.close()

    if not kerekek:
        raise EllenorzesiHiba("Nincs fogaskerék adatsor a munkalapon")
    return kerekek


def json_kiiras(kerekek, utvonal):
    """A fogaskerekeket (a számított méretekkel) JSON fájlba írja."""
    utvonal = Path(utvonal)
    utvonal.parent.mkdir(parents=True, exist_ok=True)
    with utvonal.open("w", encoding="utf-8") as fajl:
        json.dump([k.szotarba() for k in kerekek], fajl, indent=2, ensure_ascii=False)


def fo(argumentumok):
    utvonal = Path(argumentumok[1]) if len(argumentumok) > 1 else Path(
        "adat/fogaskerekek.xlsx"
    )
    try:
        kerekek = fogaskerekek_beolvasasa(utvonal)
    except (EllenorzesiHiba, FileNotFoundError) as hiba:
        print(f"Hiba: {hiba}")
        return 1

    for kerek in kerekek:
        print(
            f"{kerek.nev}: m={kerek.modul}, z={kerek.fogszam}, "
            f"d={kerek.osztokor_atmero:.2f} mm, "
            f"furat={kerek.furat_atmero:.4f} mm ({kerek.turesmod})"
        )
    json_kiiras(kerekek, "kimenet/fogaskerekek.json")
    print(f"{len(kerekek)} fogaskerék beolvasva, mentve: kimenet/fogaskerekek.json")
    return 0


if __name__ == "__main__":
    sys.exit(fo(sys.argv))
